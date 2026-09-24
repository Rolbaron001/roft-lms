#!/usr/bin/env python3
"""Turn one of this project's markdown notes into a Word document.

    python scripts/build-note-docx.py FOR-HEIDI-process-flow-and-costing.md
    python scripts/build-note-docx.py <note.md> --out "../../Design/Server/Note.docx"
    python scripts/build-note-docx.py <note.md> --style plain

Curiosa circulate these to the QCTO and to their own clients, so the markdown
in the repository is the working copy and the Word file is what leaves the
building. Writing the Word file by hand meant the two drifted, which is how a
job sheet came to disagree with itself.

Two styles:

  roft    (default) The ROFT documents already issued, as in
          scripts/build-dictionary-docx.py: Montserrat headings on navy, Inter
          body, gold accent. For notes that go out under ROFT's name.
  plain   Unbranded but finished: a title band, the For / From / Date lines set
          out as a small table, Calibri throughout, a quiet teal accent, and
          page numbers. Roland, 24 September: a note to a supplier "doesn't need
          my branding, but give it some style".

It handles the subset of markdown these notes actually use: headings,
paragraphs, bullet lists, tables, block quotes, fenced code, horizontal rules,
markdown hard breaks, and inline bold, italic and code. Anything else is
written through as plain text rather than silently dropped.
"""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


@dataclass(frozen=True)
class Style:
    heading_font: str
    body_font: str
    ink: str          # body text
    heading: str      # headings
    accent: str       # rules, quote bars, the title band
    header_fill: str  # table header row
    header_ink: str   # table header text
    stripe: str       # alternate table rows
    title_band: bool  # the title set in a coloured band
    meta_table: bool  # the For / From / Date lines as a table
    page_numbers: bool


STYLES = {
    "roft": Style(
        heading_font="Montserrat",
        body_font="Inter",
        ink="1A1A1A",
        heading="0D1E32",
        accent="B9975B",
        header_fill="0D1E32",
        header_ink="FFFFFF",
        stripe="F4F1EA",
        title_band=False,
        meta_table=False,
        page_numbers=False,
    ),
    "plain": Style(
        heading_font="Calibri",
        body_font="Calibri",
        ink="1F2933",
        heading="1D5C63",
        accent="1D5C63",
        header_fill="1D5C63",
        header_ink="FFFFFF",
        stripe="EEF4F5",
        title_band=True,
        meta_table=True,
        page_numbers=True,
    ),
}

MONO_FONT = "Consolas"

# How a hard break travels from the markdown reader to the renderer.
#
# Markdown writes one as two trailing spaces, which a naive join would throw
# away along with the "For / From / Date" block at the head of every one of
# these notes. A vertical tab is used because it cannot occur in the source.
#
# Written as the escape "\v", never as the character itself. It was once the
# raw control character, which displays as nothing at all: the line read
# `BREAK = ""`, indistinguishable from the empty string that had corrupted this
# file hours earlier. See .claude/skills/editing-files.
BREAK = "\v"

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*|" + re.escape(BREAK) + r")", re.S)
META_LINE = re.compile(r"^\*\*(.+?):\*\*\s*(.*)$")


# ---------------------------------------------------------------- docx helpers

def shade(cell, hex_colour: str) -> None:
    """Fill a table cell. python-docx has no API for this."""
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(element)


def rule(paragraph, colour: str, edge: str = "bottom", size: int = 6) -> None:
    """A horizontal line on one edge of a paragraph."""
    borders = OxmlElement("w:pBdr")
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), "8")
    element.set(qn("w:color"), colour)
    borders.append(element)
    paragraph._p.get_or_add_pPr().append(borders)


def no_borders(table) -> None:
    """Strip every border from a table, for layout rather than data."""
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    properties.append(borders)


# The page is A4 with 2.4 cm margins either side.
USABLE_CM = 21.0 - 2 * 2.4


def set_widths(table, widths_cm: list[float]) -> None:
    """Fix each column's width, in both places Word and LibreOffice read it.

    python-docx leaves every column equal and lets the reader autofit, and
    LibreOffice ignores a width set on the cells alone. Found by rendering the
    first note to InspireTec on 24 September: a label column took half the
    page, and a column holding a single digit was as wide as a sentence.
    """
    table.autofit = False
    grid = table._tbl.tblGrid
    for column, width in zip(grid.findall(qn("w:gridCol")), widths_cm):
        column.set(qn("w:w"), str(int(width * 567)))  # twentieths of a point
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def proportional_widths(header: list[str], rows: list[list[str]]) -> list[float]:
    """Widths that follow what each column holds, within sensible bounds."""
    def length(text: str) -> int:
        return len(re.sub(r"[*`]", "", text))

    longest = [
        max([length(header[i])] + [length(row[i]) for row in rows if i < len(row)])
        for i in range(len(header))
    ]
    # Short columns get enough room for their content and no more; the rest is
    # shared among the long ones by how much each has to say.
    weights = [min(max(n, 3), 60) for n in longest]
    total = sum(weights)
    widths = [max(USABLE_CM * w / total, 1.1) for w in weights]
    scale = USABLE_CM / sum(widths)
    return [w * scale for w in widths]


def keep_together(table, repeat_header: bool = True) -> None:
    """Keep a short table on one page, and repeat a long one's header row."""
    rows = table.rows
    if repeat_header and rows:
        properties = rows[0]._tr.get_or_add_trPr()
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        properties.append(header)
    # A table of a dozen rows or fewer is kept whole. A longer one may break,
    # with its header repeated, rather than leaving a gap of most of a page.
    if len(rows) <= 12:
        for row in rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True


def field(paragraph, instruction: str) -> None:
    """A Word field such as PAGE, which Word fills in when it lays out."""
    simple = OxmlElement("w:fldSimple")
    simple.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    simple.append(run)
    paragraph._p.append(simple)


def write_inline(paragraph, text: str, style: Style, *, size: float = 10.5,
                 colour: str | None = None, bold: bool = False,
                 font: str | None = None) -> None:
    """Add text to a paragraph, honouring **bold**, *italic*, `code` and breaks."""
    colour = colour or style.ink
    font = font or style.body_font
    for piece in INLINE.split(text):
        if not piece:
            continue

        if piece == BREAK:
            paragraph.add_run().add_break()
            continue

        run = paragraph.add_run()
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(colour)
        run.bold = bold
        run.font.name = font

        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            run.text = piece[2:-2]
            run.bold = True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            run.text = piece[1:-1]
            run.font.name = MONO_FONT
            run.font.size = Pt(size - 1)
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            run.text = piece[1:-1]
            run.italic = True
        else:
            run.text = piece


# -------------------------------------------------------------- markdown model

def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_divider(line: str) -> bool:
    return bool(re.fullmatch(r"\|?[\s:|-]+\|?", line.strip())) and "-" in line


def ends_paragraph(line: str) -> bool:
    """Whether this line starts a different block and so closes a paragraph."""
    return bool(re.match(r"(#{1,4}\s|[-*]\s|\d+\.\s|\||>|```|-{3,}$)", line.strip()))


def with_break(line: str) -> str:
    """The line's text, carrying a hard break if markdown asked for one."""
    return line.strip() + (BREAK if line.endswith("  ") else "")


def blocks(lines: list[str]):
    """Walk the markdown once, yielding (kind, payload)."""
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            body = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                body.append(lines[i])
                i += 1
            i += 1
            yield "code", body
            continue

        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            yield "rule", None
            i += 1
            continue

        heading = re.match(r"(#{1,4})\s+(.*)", stripped)
        if heading:
            yield f"h{len(heading.group(1))}", heading.group(2)
            i += 1
            continue

        # A table: a pipe row followed by a divider row.
        if stripped.startswith("|") and i + 1 < len(lines) and is_divider(lines[i + 1]):
            header = split_row(stripped)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i].strip()))
                i += 1
            yield "table", (header, rows)
            continue

        if stripped.startswith(">"):
            body = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                body.append(lines[i].strip().lstrip(">").strip())
                i += 1
            yield "quote", " ".join(body)
            continue

        bullet = re.match(r"[-*]\s+(.*)", stripped)
        if bullet:
            body = [bullet.group(1)]
            i += 1
            # Continuation lines are indented under the bullet.
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip() \
                    and not re.match(r"[-*]\s+", lines[i].strip()):
                body.append(lines[i].strip())
                i += 1
            yield "bullet", " ".join(body)
            continue

        numbered = re.match(r"(\d+)\.\s+(.*)", stripped)
        if numbered:
            body = [numbered.group(2)]
            i += 1
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip() \
                    and not re.match(r"(\d+\.|[-*])\s+", lines[i].strip()):
                body.append(lines[i].strip())
                i += 1
            yield "numbered", " ".join(body)
            continue

        # A paragraph runs until a blank line or the start of another block.
        body = [with_break(line)]
        i += 1
        while i < len(lines) and lines[i].strip() and not ends_paragraph(lines[i]):
            body.append(with_break(lines[i]))
            i += 1
        # Joined with spaces, except where a hard break already ended a line.
        yield "paragraph", " ".join(body).replace(BREAK + " ", BREAK)


def meta_rows(paragraph: str) -> list[tuple[str, str]] | None:
    """The For / From / Date lines, if that is what this paragraph is."""
    lines = [line.strip() for line in paragraph.split(BREAK) if line.strip()]
    rows = [META_LINE.match(line) for line in lines]
    if len(lines) >= 2 and all(rows):
        return [(match.group(1), match.group(2)) for match in rows]
    return None


# ------------------------------------------------------------------- rendering

def title_band(document, text: str, style: Style) -> None:
    """The document's title, set in a band of the accent colour."""
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(table)
    cell = table.rows[0].cells[0]
    shade(cell, style.accent)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.left_indent = Cm(0.3)
    write_inline(paragraph, text, style, size=20, colour="FFFFFF", bold=True,
                 font=style.heading_font)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def meta_table(document, rows: list[tuple[str, str]], style: Style) -> None:
    """For / From / Date as a quiet two-column table under the title."""
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    no_borders(table)
    for label, value in rows:
        cells = table.add_row().cells
        for cell in cells:
            shade(cell, style.stripe)
        write_inline(cells[0].paragraphs[0], label, style, size=9.5,
                     colour=style.heading, bold=True)
        write_inline(cells[1].paragraphs[0], value, style, size=9.5)
    set_widths(table, [2.6, USABLE_CM - 2.6])
    keep_together(table, repeat_header=False)
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def build(markdown: str, out: Path, style: Style) -> None:
    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = style.body_font
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(style.ink)

    for section in document.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)

        if style.page_numbers:
            footer = section.footer.paragraphs[0]
            footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            write_inline(footer, "Page ", style, size=8.5, colour="6B7680")
            field(footer, "PAGE")
            write_inline(footer, " of ", style, size=8.5, colour="6B7680")
            field(footer, "NUMPAGES")

    lines = markdown.replace("\r\n", "\n").split("\n")
    first_heading = True
    after_title = False

    for kind, payload in blocks(lines):
        if kind == "h1":
            if style.title_band:
                title_band(document, payload, style)
            else:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(4)
                write_inline(paragraph, payload, style, size=22,
                             colour=style.heading, bold=True,
                             font=style.heading_font)
                rule(paragraph, style.accent, size=12)
            after_title = True
            continue

        # The block straight after the title, if it is the For / From lines.
        if after_title and kind == "paragraph" and style.meta_table:
            rows = meta_rows(payload)
            if rows:
                meta_table(document, rows, style)
                after_title = False
                continue
        after_title = False

        if kind in ("h2", "h3", "h4"):
            size = {"h2": 14, "h3": 11.5, "h4": 10.5}[kind]
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(14 if kind == "h2" else 10)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_with_next = True
            write_inline(paragraph, payload, style, size=size,
                         colour=style.heading, bold=True,
                         font=style.heading_font)
            if kind == "h2" and not first_heading:
                rule(paragraph, style.accent, size=4)
            first_heading = False

        elif kind == "paragraph":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(7)
            write_inline(paragraph, payload, style)

        elif kind == "bullet":
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(3)
            write_inline(paragraph, payload, style)

        elif kind == "numbered":
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.space_after = Pt(3)
            write_inline(paragraph, payload, style)

        elif kind == "quote":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(9)
            write_inline(paragraph, payload, style, size=11, colour=style.heading)
            rule(paragraph, style.accent, edge="left", size=18)

        elif kind == "code":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(9)
            for index, line in enumerate(payload):
                run = paragraph.add_run(line)
                run.font.name = MONO_FONT
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor.from_string(style.heading)
                if index < len(payload) - 1:
                    run.add_break()

        elif kind == "rule":
            # In the plain style every section heading carries its own rule, so
            # a `---` between sections drew a second line beside it, or a
            # stray one at the foot of a page. It becomes a little space.
            if style.title_band:
                document.add_paragraph().paragraph_format.space_after = Pt(2)
            else:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(2)
                rule(paragraph, style.accent, size=4)

        elif kind == "table":
            header, rows = payload
            # The sentence introducing a table goes to the next page with it,
            # rather than being left at the foot of one page pointing at a
            # table that starts the next.
            if document.paragraphs:
                document.paragraphs[-1].paragraph_format.keep_with_next = True
            table = document.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            for cell, text in zip(table.rows[0].cells, header):
                shade(cell, style.header_fill)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(2)
                write_inline(paragraph, text.replace("**", ""), style, size=9.5,
                             colour=style.header_ink, bold=True,
                             font=style.heading_font)

            for index, row in enumerate(rows):
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    if index % 2 == 1:
                        shade(cell, style.stripe)
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    write_inline(paragraph, text, style, size=9.5)

            set_widths(table, proportional_widths(header, rows))
            keep_together(table)
            document.add_paragraph().paragraph_format.space_after = Pt(4)

    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(out)
    print(f"  {out}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("source", type=Path, help="the markdown note")
    parser.add_argument("--out", type=Path, default=None,
                        help="where the Word file goes (default: beside the note)")
    parser.add_argument("--style", choices=sorted(STYLES), default="roft",
                        help="roft for notes under ROFT's name, plain for everything else")
    args = parser.parse_args()

    source: Path = args.source
    if not source.exists():
        raise SystemExit(f"  {source} is not there.")

    build(source.read_text(encoding="utf-8"),
          args.out or source.with_suffix(".docx"),
          STYLES[args.style])


if __name__ == "__main__":
    main()
